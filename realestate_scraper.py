"""Real Estate Listings Scraper: paginated property HTML to formatted Excel.

Use --demo for a reproducible HTML-fixture run, --demo-browser to also exercise
JavaScript rendering in Chromium, or --serve for the local portfolio dashboard.
Live selectors target Realtor.com-style markup; external compatibility varies.
"""

import argparse
import datetime
import io
import json
import logging
import math
import os
import random
import re
import tempfile
import time
from contextlib import suppress
from dataclasses import asdict, dataclass, field
from html import escape
from pathlib import Path
from urllib.parse import parse_qsl, quote, urlencode, urljoin, urlsplit, urlunsplit

import pandas as pd
from bs4 import BeautifulSoup
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import TimeoutError as PlaywrightTimeout
from playwright.sync_api import sync_playwright

logger = logging.getLogger(__name__)
CARD = '[data-testid="card-content"]'
EMPTY = '[data-testid="no-results"]'
NEXT = 'a[rel="next"], a[data-testid="pagination-next"]'


@dataclass
class Listing:
    address: str = ""
    price: str = ""
    beds: str = ""
    baths: str = ""
    sqft: str = ""
    listing_url: str = ""
    date_scraped: str = field(default_factory=lambda: datetime.date.today().isoformat())


def number(value: str, kind: str = "number") -> float | None:
    """Parse a single scalar; never turn unknown values/ranges into zero."""
    text = str(value).strip().lower().replace(",", "")
    if text == "studio" and kind == "beds":
        return 0
    units = {
        "price": r"(?:usd)?",
        "beds": r"(?:beds?|bedrooms?|bd)?",
        "baths": r"(?:baths?|bathrooms?|ba)?",
        "sqft": r"(?:sq\s*\.?\s*ft\.?|sqft|ft²)?",
    }
    match = re.fullmatch(r"\$?\s*(\d+(?:\.\d+)?)\s*([km]?)\s*" + units.get(kind, ""), text)
    if not match:
        return None
    result = float(match[1]) * {"": 1, "k": 1000, "m": 1000000}[match[2]]
    if not math.isfinite(result):
        return None
    if kind in {"beds", "sqft"} and not result.is_integer():
        return None
    return result


def safe_url(href: str, base: str = "https://www.realtor.com/") -> str:
    if not href or not href.strip() or href.strip().startswith("#"):
        return ""
    try:
        parts = urlsplit(urljoin(base, href.strip()))
        if (
            parts.scheme not in {"https", "http"}
            or not parts.hostname
            or parts.username
            or parts.password
        ):
            return ""
        query = [
            (k, v)
            for k, v in parse_qsl(parts.query, keep_blank_values=True)
            if not k.lower().startswith("utm_") and k.lower() not in {"fbclid", "gclid"}
        ]
        return urlunsplit(parts._replace(query=urlencode(query), fragment=""))
    except ValueError:
        return ""


class PageParseError(RuntimeError):
    """A failed/unknown page must not be reported as a successful empty search."""


class RealEstateScraper:
    BASE_URL = "https://www.realtor.com/realestateandhomes-search"

    def __init__(self, city: str, max_price: int, min_beds: int, headless: bool = True):
        self.city = city.strip()
        if not self.city or not re.search(r"\w", self.city):
            raise ValueError("City must not be blank")
        if max_price <= 0 or min_beds < 0:
            raise ValueError("Maximum price must be positive and minimum beds cannot be negative")
        self.max_price, self.min_beds, self.headless = max_price, min_beds, headless
        self.last_error = None
        self.stats = {}

    def _build_url(self, page_num: int = 1) -> str:
        if page_num < 1:
            raise ValueError("Page number must be positive")
        city_slug = quote(re.sub(r"[\s,]+", "_", self.city).strip("_"), safe="_-")
        suffix = f"/pg-{page_num}" if page_num > 1 else ""
        return (
            f"{self.BASE_URL}/{city_slug}/price-na-{self.max_price}/beds-{self.min_beds}-na{suffix}"
        )

    def _parse_card(self, card, page=None) -> Listing | None:
        def element(selector):
            try:
                return card.query_selector(selector)
            except Exception:
                return None

        def text(selector):
            try:
                el = element(selector)
                return re.sub(r"\s+", " ", el.inner_text()).strip() if el else ""
            except Exception:
                return ""

        price = text('[data-testid="card-price"]') or text(
            ".Pricestyles__Component-rui__sc-rmiu7q-0"
        )
        address = ", ".join(
            filter(
                None,
                [text('[data-testid="card-address-1"]'), text('[data-testid="card-address-2"]')],
            )
        )
        fields = {
            key: text(f'[data-testid="property-meta-{key}"]') for key in ("beds", "baths", "sqft")
        }
        meta = text(".property-meta")
        patterns = {
            "beds": r"(\d+)\s*(?:bed|bd)",
            "baths": r"(\d+(?:\.\d+)?)\s*(?:bath|ba)",
            "sqft": r"([\d,]+)\s*(?:sq\s*\.?\s*ft|ft²)",
        }
        for key, pattern in patterns.items():
            if not fields[key]:
                match = re.search(pattern, meta, re.I)
                fields[key] = (
                    match[1]
                    if match
                    else ("0" if key == "beds" and "studio" in meta.lower() else "")
                )
            parsed = number(fields[key], key)
            if parsed is not None:
                fields[key] = f"{int(parsed):,}" if key == "sqft" else f"{parsed:g}"
        url = ""
        # Prefer a property link over unrelated links such as broker profiles.
        link = element('a[href*="realestateandhomes-detail"]') or element("a")
        if link:
            with suppress(Exception):
                base = getattr(page, "url", "") or "https://www.realtor.com/"
                if not base.startswith("http"):
                    base = "https://www.realtor.com/"
                url = safe_url(link.get_attribute("href") or "", base)
        if not address and not price:
            return None
        return Listing(address=address, price=price, listing_url=url, **fields)

    def _parse_listings(self, page) -> list[Listing]:
        try:
            page.wait_for_selector(f"{CARD}, {EMPTY}", timeout=15000)
        except PlaywrightTimeout as exc:
            raise PageParseError("No recognized listings or empty-results marker appeared") from exc
        cards = page.query_selector_all(CARD)
        result = []
        for card in cards:
            listing = self._parse_card(card, page)
            if listing:
                result.append(listing)
        if cards and not result:
            raise PageParseError("Listing cards appeared but none could be parsed")
        return result

    def _collect(self, loader, max_pages, throttle):
        self.last_error = None
        self.stats = {"pages": 0, "duplicates": 0, "excluded": 0, "events": []}
        results, seen, visited = [], set(), set()
        url = self._build_url()
        for page_num in range(1, max_pages + 1):
            visited.add(url)
            logger.info("Loading page %s: %s", page_num, url)
            try:
                page = loader(url)
                if page is None:
                    raise PageParseError("Page did not load")
                listings = self._parse_listings(page)
            except (PlaywrightError, PageParseError, OSError) as exc:
                self.last_error = f"Page {page_num} failed: {exc}"
                logger.warning(self.last_error)
                break
            self.stats["pages"] += 1
            added = 0
            for listing in listings:
                price, beds = number(listing.price, "price"), number(listing.beds, "beds")
                if price is None or beds is None or price > self.max_price or beds < self.min_beds:
                    self.stats["excluded"] += 1
                    continue
                key = listing.listing_url or re.sub(r"\s+", " ", listing.address).casefold()
                if key and key in seen:
                    self.stats["duplicates"] += 1
                    continue
                if key:
                    seen.add(key)
                results.append(listing)
                added += 1
            self.stats["events"].append({"page": page_num, "parsed": len(listings), "added": added})
            if not listings:
                break
            try:
                next_link = page.query_selector(NEXT)
                if (
                    not next_link
                    or next_link.get_attribute("aria-disabled") == "true"
                    or next_link.get_attribute("disabled") is not None
                ):
                    break
                next_url = safe_url(next_link.get_attribute("href") or "", url)
            except PlaywrightError as exc:
                self.last_error = f"Pagination on page {page_num} failed: {exc}"
                break
            if not next_url or urlsplit(next_url).netloc != urlsplit(self.BASE_URL).netloc:
                self.last_error = "Invalid or off-site pagination link; stopped"
                break
            if next_url in visited:
                self.last_error = "Repeated pagination URL; stopped to prevent a loop"
                break
            url = next_url
            if throttle and page_num < max_pages:
                time.sleep(random.uniform(4, 8))
        logger.info("Collected %s unique matching listings", len(results))
        return results

    def scrape(self, max_pages: int = 5, *, page_loader=None, throttle=True) -> list[Listing]:
        if max_pages < 1:
            raise ValueError("Maximum pages must be positive")
        if page_loader is not None:
            return self._collect(page_loader, max_pages, throttle)
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=self.headless)
            try:
                context = browser.new_context(viewport={"width": 1280, "height": 900})
                page = context.new_page()

                def load(url):
                    response = page.goto(url, wait_until="domcontentloaded", timeout=30000)
                    if response and response.status >= 400:
                        raise PageParseError(f"HTTP {response.status}")
                    return page

                return self._collect(load, max_pages, throttle)
            finally:
                browser.close()


class HTMLFixture:
    """The selector/text interface shared by saved HTML and Playwright elements."""

    def __init__(self, element, url="https://www.realtor.com/"):
        self.element = BeautifulSoup(element, "lxml") if isinstance(element, str) else element
        self.url = url

    def query_selector(self, selector):
        found = self.element.select_one(selector)
        return HTMLFixture(found, self.url) if found is not None else None

    def query_selector_all(self, selector):
        return [HTMLFixture(el, self.url) for el in self.element.select(selector)]

    def inner_text(self):
        return self.element.get_text(" ", strip=True)

    def get_attribute(self, name):
        return self.element.get(name)

    def wait_for_selector(self, selector, timeout=15000):
        if not self.query_selector(selector):
            raise PlaywrightTimeout("Fixture has no recognized results")


def generate_demo_listings(
    city: str, max_price: int, min_beds: int, *, count=25, seed=42
) -> list[Listing]:
    """Reproducible synthetic records, never represented as live property listings."""
    RealEstateScraper(city, max_price, min_beds)
    if not 1 <= count <= 500:
        raise ValueError("Demo count must be 1–500")
    rng = random.Random(seed)
    streets = [
        "Oak Ridge Drive",
        "Maple Avenue",
        "Sunset Lane",
        "Cedar Hollow",
        "Willow Creek Way",
        "Riverstone Parkway",
        "Hillcrest Drive",
        "Magnolia Court",
        "Lake View Lane",
        "Canyon Road",
    ]
    zipcode = {"austin": "78704", "denver": "80206", "raleigh": "27608"}.get(
        city.split(",")[0].strip().lower(), ""
    )
    records = []
    for i in range(count):
        beds = rng.randint(min_beds, min_beds + 2)
        price = rng.randint(max(1, int(max_price * 0.55)), max(1, int(max_price * 0.97)))
        records.append(
            Listing(
                address=f"{120 + 43 * i} {streets[i % len(streets)]}, {city.strip()} {zipcode}".strip(),
                price=f"${price:,}",
                beds=str(beds),
                baths=f"{rng.choice([1.5, 2, 2.5, 3, 3.5]):g}",
                sqft=f"{rng.randint(1200, 3400):,}",
                listing_url=f"https://homes.example/listings/{i + 1:03d}",
            )
        )
    return records


def demo_pages(city, max_price, min_beds, count=25, seed=42):
    records = generate_demo_listings(city, max_price, min_beds, count=count, seed=seed)
    scraper = RealEstateScraper(city, max_price, min_beds)

    def card(listing):
        parts = listing.address.split(",", 1)
        meta = f'<div class="property-meta">{listing.beds} bed {listing.baths} bath {listing.sqft} sq ft</div>'
        return (
            f'<article data-testid="card-content"><a href="{escape(listing.listing_url)}">Details</a>'
            f'<div data-testid="card-price">{escape(listing.price)}</div>'
            f'<div data-testid="card-address-1">{escape(parts[0])}</div>'
            f'<div data-testid="card-address-2">{escape(parts[1].strip()) if len(parts) > 1 else ""}</div>{meta}</article>'
        )

    pages, page_count = {}, (count + 8) // 9
    for index in range(page_count):
        chunk = records[index * 9 : (index + 1) * 9]
        if index:
            chunk = [records[index * 9 - 1], *chunk]
        html = "".join(card(record) for record in chunk)
        if index + 1 < page_count:
            html += f'<a rel="next" href="{escape(scraper._build_url(index + 2))}">Next</a>'
        pages[scraper._build_url(index + 1)] = html
    return scraper, pages


def run_demo(city="Austin, TX", max_price=500000, min_beds=3, *, count=25, seed=42, browser=False):
    scraper, pages = demo_pages(city, max_price, min_beds, count, seed)
    if browser:
        with sync_playwright() as playwright:
            chromium = playwright.chromium.launch(headless=True)
            try:
                page = chromium.new_page()
                page.route("**/*", lambda route: route.abort())

                def load(url):
                    # All markup is generated locally and inserted asynchronously.
                    script = json.dumps(pages[url]).replace("<", "\\u003c")
                    page.set_content(
                        f'<main id="results"></main><script>setTimeout(()=>{{document.getElementById("results").innerHTML={script}}},50)</script>'
                    )
                    return page

                records = scraper.scrape(len(pages), page_loader=load, throttle=False)
            finally:
                chromium.close()
    else:
        records = scraper.scrape(
            len(pages), page_loader=lambda url: HTMLFixture(pages[url], url), throttle=False
        )
    stats = {
        **scraper.stats,
        "seed": seed,
        "mode": "Chromium · JavaScript fixture" if browser else "HTML fixture",
        "source": "Synthetic properties • generated HTML • no external requests",
    }
    if scraper.last_error:
        raise PageParseError(scraper.last_error)
    return records, stats


HEADERS = ["Address", "Price", "Beds", "Baths", "Sq Ft", "Listing URL", "Date Scraped"]


def export_to_excel(listings: list[Listing], output_path, *, source="Property listings") -> None:
    if not listings:
        logger.warning("No listings to export")
        return
    rows = []
    for listing in listings:
        row = {k: ILLEGAL_CHARACTERS_RE.sub("", str(v)) for k, v in asdict(listing).items()}
        for key in ("price", "beds", "baths", "sqft"):
            parsed = number(row[key], key)
            if parsed is not None:
                row[key] = parsed
        with suppress(ValueError):
            row["date_scraped"] = datetime.date.fromisoformat(row["date_scraped"])
        rows.append(row)
    frame = pd.DataFrame(rows)
    frame.columns = HEADERS
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl", date_format="yyyy-mm-dd") as writer:
        frame.to_excel(writer, index=False, sheet_name="Listings")
        ws, wb = writer.sheets["Listings"], writer.book
        wb.properties.title = "Habitat | Property Listings"
        wb.properties.description = source
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        for row in ws:
            ws.row_dimensions[row[0].row].height = 29
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11, color="243448")
                cell.alignment = Alignment(vertical="center")
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F3F3ED")
            if row[0].row > 1:
                row[1].number_format = '"$"#,##0'
                row[2].number_format = "0"
                row[3].number_format = "0.0"
                row[4].number_format = "#,##0"
                row[6].number_format = "yyyy-mm-dd"
        for cell in ws[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="223448")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                max(len(str(c.value or "")) for c in col) + 4, 64
            )
        notes = wb.create_sheet("Export notes")
        for row in [
            ["HABITAT", "PROPERTY EXPORT"],
            ["Source", source],
            ["Listings", len(listings)],
            ["Missing data", "Blank or unparsed values are not assumed to be zero."],
            [
                "Synthetic data",
                "Demo addresses and prices are fictional; listing links use .example.",
            ],
            ["Price", "Asking price from the source; not a valuation."],
        ]:
            notes.append(row)
        notes.column_dimensions["A"].width = 22
        notes.column_dimensions["B"].width = 92
        notes.sheet_view.showGridLines = False
        for row in notes:
            notes.row_dimensions[row[0].row].height = 30
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11, color="243448")
        for cell in notes[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="223448")
    if hasattr(output_path, "write"):
        output_path.write(buffer.getvalue())
    else:
        path = Path(output_path)
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Output filename must end in .xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)
        temp = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=path.parent, suffix=".xlsx", delete=False
            ) as handle:
                temp = Path(handle.name)
                handle.write(buffer.getvalue())
            os.replace(temp, path)
        finally:
            if temp and temp.exists():
                temp.unlink()
    logger.info("Exported %s listings", len(listings))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--city", default="Austin, TX")
    parser.add_argument("--max-price", type=int, default=500000)
    parser.add_argument("--min-beds", type=int, default=3)
    parser.add_argument("--max-pages", type=int, default=3)
    parser.add_argument("--output", default=None)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--demo", action="store_true")
    mode.add_argument(
        "--demo-browser",
        action="store_true",
        help="Parse locally rendered JavaScript fixtures in Chromium",
    )
    mode.add_argument("--serve", action="store_true", help="Start the local offline dashboard")
    parser.add_argument("--port", type=int, default=8766)
    parser.add_argument("--count", type=int, default=25)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--no-headless", action="store_true")
    args = parser.parse_args(argv)
    try:
        scraper = RealEstateScraper(args.city, args.max_price, args.min_beds, not args.no_headless)
        if args.max_pages < 1 or not 1 <= args.count <= 500 or not 1 <= args.port <= 65535:
            raise ValueError("Pages must be positive, demo count 1–500, and port 1–65535")
        if args.output and Path(args.output).suffix.lower() != ".xlsx":
            raise ValueError("Output filename must end in .xlsx")
    except ValueError as exc:
        parser.error(str(exc))
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    try:
        if args.serve:
            from dashboard import serve

            serve(args.port)
            return 0
        demo = args.demo or args.demo_browser
        if demo:
            records, stats = run_demo(
                args.city,
                args.max_price,
                args.min_beds,
                count=args.count,
                seed=args.seed,
                browser=args.demo_browser,
            )
            source = stats["source"] + " | " + stats["mode"]
            logger.info(
                "%s: %s pages, %s duplicates removed",
                stats["mode"],
                stats["pages"],
                stats["duplicates"],
            )
        else:
            records = scraper.scrape(args.max_pages)
            source = f"Realtor.com | {args.city}" + (
                f" | PARTIAL: {scraper.last_error}" if scraper.last_error else ""
            )
        if not records:
            logger.error(scraper.last_error or "No matching listings found")
            return 1
        output = args.output or ("listings_demo.xlsx" if demo else "listings.xlsx")
        export_to_excel(records, output, source=source)
        if scraper.last_error:
            print(f"[PARTIAL] {len(records)} listings saved to {output}. {scraper.last_error}")
            return 2
        print(f"[OK] {len(records)} listings saved to: {output}")
        return 0
    except (OSError, ValueError, PlaywrightError, PageParseError) as exc:
        logger.error("%s", exc)
        return 1
    except KeyboardInterrupt:
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
