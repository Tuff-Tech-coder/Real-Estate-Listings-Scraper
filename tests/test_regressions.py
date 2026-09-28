"""Cases reproduced against the original version before repair."""

import importlib

import openpyxl
import pytest
from test_realestate_scraper import FULL, FakeCard, FakeElement

import realestate_scraper as rs


def scraper():
    return rs.RealEstateScraper("Austin, TX", 500000, 3)


def test_partial_meta_fallback():
    mapping = {k: v for k, v in FULL.items() if "baths" not in k and "sqft" not in k}
    mapping[".property-meta"] = "3 bed 2.5 bath 1,850 sq ft"
    listing = scraper()._parse_card(FakeCard(mapping), None)
    assert listing.baths == "2.5"
    assert listing.sqft == "1,850"


def test_fallback_preserves_existing_values():
    mapping = {k: v for k, v in FULL.items() if "beds" not in k}
    mapping[".property-meta"] = "3 bed"
    listing = scraper()._parse_card(FakeCard(mapping), None)
    assert listing.baths == "2"
    assert listing.sqft == "1,847"


def test_second_address_line_without_first_has_no_leading_comma():
    mapping = {k: v for k, v in FULL.items() if "address-1" not in k}
    assert scraper()._parse_card(FakeCard(mapping), None).address == "Austin, TX 78704"


def test_protocol_relative_listing_link():
    card = FakeCard(FULL, link_href="//www.realtor.com/detail/123")
    assert scraper()._parse_card(card, None).listing_url == "https://www.realtor.com/detail/123"


def test_empty_href_stays_empty():
    card = FakeCard(FULL)
    card._link = FakeElement(attrs={"href": ""})
    assert scraper()._parse_card(card, None).listing_url == ""


def test_unsafe_link_is_not_a_listing_url():
    card = FakeCard(FULL, link_href="javascript:alert(1)")
    assert scraper()._parse_card(card, None).listing_url == ""


def test_detached_link_does_not_discard_property():
    listing = scraper()._parse_card(FakeCard(FULL, raises_on=["a"]), None)
    assert listing.price == "$342,500"
    assert listing.listing_url == ""


def test_city_without_space_after_comma():
    assert "/Austin_TX/" in rs.RealEstateScraper("Austin,TX", 500000, 3)._build_url()


def test_seven_bedroom_demo_does_not_crash():
    assert all(int(x.beds) >= 7 for x in rs.generate_demo_listings("Austin, TX", 900000, 7))


def test_demo_uses_city_appropriate_zip():
    assert all("787" not in x.address for x in rs.generate_demo_listings("Denver, CO", 500000, 3))


def test_demo_links_do_not_impersonate_live_listings():
    assert all(
        ".example/" in x.listing_url for x in rs.generate_demo_listings("Austin, TX", 500000, 3)
    )


def test_nested_export_folder(tmp_path):
    out = tmp_path / "exports" / "listings.xlsx"
    rs.export_to_excel([rs.Listing(address="Example")], out)
    assert out.exists()


def test_export_price_is_numeric(tmp_path):
    out = tmp_path / "listings.xlsx"
    rs.export_to_excel([rs.Listing(price="$350,000", beds="3", baths="2.5", sqft="1,500")], out)
    ws = openpyxl.load_workbook(out)["Listings"]
    assert [ws.cell(2, i).value for i in (2, 3, 4, 5)] == [350000, 3, 2.5, 1500]


def test_formula_like_address_stays_text(tmp_path):
    out = tmp_path / "listings.xlsx"
    rs.export_to_excel([rs.Listing(address="=1+1")], out)
    assert openpyxl.load_workbook(out)["Listings"]["A2"].data_type == "s"


def test_export_control_character_is_cleaned(tmp_path):
    out = tmp_path / "listings.xlsx"
    rs.export_to_excel([rs.Listing(address="1\x01 Oak Lane")], out)
    assert openpyxl.load_workbook(out)["Listings"]["A2"].value == "1 Oak Lane"


def test_import_has_no_log_side_effect(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    importlib.reload(rs)
    assert not (tmp_path / "scraper.log").exists()


def test_negative_bedroom_filter_is_rejected():
    with pytest.raises(ValueError):
        rs.RealEstateScraper("Austin, TX", 500000, -1)
