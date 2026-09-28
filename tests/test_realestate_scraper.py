"""Tests for URL construction, card parsing, demo generation and Excel export.

No browser is launched and no network request is made. Playwright's element
API is small enough to stub directly -- query_selector, inner_text and
get_attribute -- which is what makes _parse_card testable without Chromium.
"""

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from realestate_scraper import (  # noqa: E402
    Listing,
    RealEstateScraper,
    export_to_excel,
    generate_demo_listings,
)


# ---------------------------------------------------------------------------
# Playwright element stubs
# ---------------------------------------------------------------------------
class FakeElement:
    """Stands in for a Playwright ElementHandle."""

    def __init__(self, text="", attrs=None):
        self._text = text
        self._attrs = attrs or {}

    def inner_text(self):
        return self._text

    def get_attribute(self, name):
        return self._attrs.get(name)


class FakeCard:
    """A listing card whose query_selector is driven by a selector->text map."""

    def __init__(self, mapping, link_href=None, raises_on=()):
        self._mapping = mapping
        self._link = FakeElement(attrs={"href": link_href}) if link_href else None
        self._raises_on = set(raises_on)

    def query_selector(self, selector):
        if selector in self._raises_on:
            raise RuntimeError("detached element")
        if selector == "a":
            return self._link
        text = self._mapping.get(selector)
        return FakeElement(text) if text is not None else None


FULL = {
    '[data-testid="card-price"]': "$342,500",
    '[data-testid="card-address-1"]': "4823 Oak Ridge Dr",
    '[data-testid="card-address-2"]': "Austin, TX 78704",
    '[data-testid="property-meta-beds"]': "3",
    '[data-testid="property-meta-baths"]': "2",
    '[data-testid="property-meta-sqft"]': "1,847",
}


@pytest.fixture
def scraper():
    return RealEstateScraper(city="Austin, TX", max_price=500000, min_beds=3)


# ---------------------------------------------------------------------------
# URL construction
# ---------------------------------------------------------------------------
class TestBuildUrl:
    def test_encodes_filters_into_the_path(self, scraper):
        url = scraper._build_url()
        assert "Austin_TX" in url
        assert "price-na-500000" in url
        assert "beds-3-na" in url

    def test_first_page_has_no_page_suffix(self, scraper):
        assert not scraper._build_url(1).endswith("/pg-1")

    def test_later_pages_append_page_suffix(self, scraper):
        assert scraper._build_url(3).endswith("/pg-3")

    def test_multiword_city_is_slugified(self):
        s = RealEstateScraper(city="San Antonio, TX", max_price=1, min_beds=1)
        assert "San_Antonio_TX" in s._build_url()

    def test_city_without_state_still_builds(self):
        s = RealEstateScraper(city="Austin", max_price=1, min_beds=1)
        assert "/Austin/" in s._build_url()


# ---------------------------------------------------------------------------
# Card parsing -- defensive extraction
# ---------------------------------------------------------------------------
class TestParseCard:
    def test_extracts_every_field(self, scraper):
        listing = scraper._parse_card(
            FakeCard(FULL, link_href="/realestateandhomes-detail/123"), None
        )
        assert listing.price == "$342,500"
        assert listing.address == "4823 Oak Ridge Dr, Austin, TX 78704"
        assert listing.beds == "3"
        assert listing.baths == "2"
        assert listing.sqft == "1,847"

    def test_relative_href_becomes_absolute(self, scraper):
        listing = scraper._parse_card(FakeCard(FULL, link_href="/detail/123"), None)
        assert listing.listing_url == "https://www.realtor.com/detail/123"

    def test_absolute_href_is_left_alone(self, scraper):
        listing = scraper._parse_card(
            FakeCard(FULL, link_href="https://www.realtor.com/detail/9"), None
        )
        assert listing.listing_url == "https://www.realtor.com/detail/9"

    def test_missing_link_yields_empty_url(self, scraper):
        assert scraper._parse_card(FakeCard(FULL), None).listing_url == ""

    def test_second_address_line_is_optional(self, scraper):
        mapping = {k: v for k, v in FULL.items() if "address-2" not in k}
        listing = scraper._parse_card(FakeCard(mapping), None)
        assert listing.address == "4823 Oak Ridge Dr"

    def test_regex_fallback_when_meta_testids_absent(self, scraper):
        """Beds/baths/sqft come from a combined string when the testids are gone."""
        mapping = {
            '[data-testid="card-price"]': "$410,000",
            '[data-testid="card-address-1"]': "12 Elm St",
            ".property-meta": "4 bed 2.5 bath 2,100 sq ft",
        }
        listing = scraper._parse_card(FakeCard(mapping), None)
        assert listing.beds == "4"
        assert listing.baths == "2.5"
        assert listing.sqft == "2,100"

    def test_regex_fallback_handles_abbreviations(self, scraper):
        mapping = {
            '[data-testid="card-address-1"]': "9 Pine Rd",
            ".property-meta": "3 bd 2 ba 1,500 sqft",
        }
        listing = scraper._parse_card(FakeCard(mapping), None)
        assert listing.beds == "3"
        assert listing.baths == "2"

    def test_price_falls_back_to_secondary_selector(self, scraper):
        mapping = {
            '[data-testid="card-address-1"]': "7 Ash Ct",
            ".Pricestyles__Component-rui__sc-rmiu7q-0": "$275,000",
        }
        assert scraper._parse_card(FakeCard(mapping), None).price == "$275,000"

    def test_card_without_address_or_price_is_rejected(self, scraper):
        assert scraper._parse_card(FakeCard({}), None) is None

    def test_price_alone_is_enough_to_keep_a_card(self, scraper):
        mapping = {'[data-testid="card-price"]': "$199,000"}
        listing = scraper._parse_card(FakeCard(mapping), None)
        assert listing is not None
        assert listing.address == ""

    def test_detached_element_does_not_raise(self, scraper):
        """safe_text swallows element errors so one bad field can't kill a card."""
        card = FakeCard(FULL, raises_on=['[data-testid="property-meta-beds"]'])
        listing = scraper._parse_card(card, None)
        assert listing is not None
        assert listing.price == "$342,500"

    def test_date_scraped_is_populated(self, scraper):
        assert scraper._parse_card(FakeCard(FULL), None).date_scraped


# ---------------------------------------------------------------------------
# Demo generation
# ---------------------------------------------------------------------------
class TestGenerateDemoListings:
    def test_produces_listings(self):
        listings = generate_demo_listings("Austin, TX", 500000, 3)
        assert len(listings) >= 20
        assert all(isinstance(x, Listing) for x in listings)

    def test_respects_the_max_price_filter(self):
        """Demo data must satisfy the filters it claims to match."""
        listings = generate_demo_listings("Austin, TX", 500000, 3)
        for listing in listings:
            value = int(listing.price.replace("$", "").replace(",", ""))
            assert value <= 500000

    def test_respects_the_min_beds_filter(self):
        listings = generate_demo_listings("Austin, TX", 500000, 3)
        assert all(int(x.beds) >= 3 for x in listings)

    def test_city_appears_in_addresses(self):
        listings = generate_demo_listings("Denver, CO", 900000, 2)
        assert all("Denver" in x.address for x in listings)


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------
class TestExportToExcel:
    @pytest.fixture
    def listings(self):
        return [
            Listing(
                address="1 Main St",
                price="$300,000",
                beds="3",
                baths="2",
                sqft="1,500",
                listing_url="https://example.com/1",
            ),
            Listing(
                address="2 Oak Ave",
                price="$425,000",
                beds="4",
                baths="3",
                sqft="2,200",
                listing_url="https://example.com/2",
            ),
        ]

    def test_writes_a_readable_workbook(self, listings, tmp_path):
        import openpyxl

        out = tmp_path / "listings.xlsx"
        export_to_excel(listings, str(out))
        ws = openpyxl.load_workbook(out)["Listings"]
        assert ws.max_row == 3
        assert ws.cell(row=2, column=1).value == "1 Main St"

    def test_headers_are_client_friendly(self, listings, tmp_path):
        import openpyxl

        out = tmp_path / "listings.xlsx"
        export_to_excel(listings, str(out))
        ws = openpyxl.load_workbook(out)["Listings"]
        assert [c.value for c in ws[1]] == [
            "Address",
            "Price",
            "Beds",
            "Baths",
            "Sq Ft",
            "Listing URL",
            "Date Scraped",
        ]

    def test_empty_input_writes_no_file(self, tmp_path):
        out = tmp_path / "empty.xlsx"
        export_to_excel([], str(out))
        assert not out.exists()
