"""End-to-end fixture, transport lifecycle, CLI, and localhost export checks."""

import io
import json
import threading
from http.client import HTTPConnection
from http.server import ThreadingHTTPServer
from types import SimpleNamespace
from unittest.mock import Mock

import openpyxl
import pytest

import realestate_scraper as rs
from dashboard import Handler, filter_records


@pytest.mark.parametrize(
    ("text", "kind", "expected"),
    [
        ("$1.2M", "price", 1200000),
        ("$350K", "price", 350000),
        ("3 beds", "beds", 3),
        ("Studio", "beds", 0),
        ("2.5 baths", "baths", 2.5),
        ("1,500 sq ft", "sqft", 1500),
        ("Contact agent", "price", None),
        ("3–4 beds", "beds", None),
        ("", "baths", None),
        ("-100", "price", None),
    ],
)
def test_number_parsing(text, kind, expected):
    assert rs.number(text, kind) == expected


def test_numeric_overflow_is_not_an_infinite_price():
    assert rs.number("9" * 400, "price") is None


def test_detached_pagination_preserves_already_parsed_records():
    scraper, pages = rs.demo_pages("Austin, TX", 500000, 3)
    fixture = rs.HTMLFixture(pages[scraper._build_url()])
    original_query = fixture.query_selector

    def query(selector):
        if selector == rs.NEXT:
            raise rs.PlaywrightError("Detached navigation")
        return original_query(selector)

    fixture.query_selector = query
    assert len(scraper.scrape(page_loader=lambda url: fixture, throttle=False)) == 9
    assert "Pagination on page 1 failed" in scraper.last_error


def test_html_demo_round_trips_records():
    records, stats = rs.run_demo()
    assert records == rs.generate_demo_listings("Austin, TX", 500000, 3)
    assert len(records) == 25
    assert stats["pages"] == 3
    assert stats["duplicates"] == 2
    assert sum(e["added"] for e in stats["events"]) == 25


def test_seed_does_not_mutate_global_random_state():
    state = rs.random.getstate()
    rs.run_demo()
    assert rs.random.getstate() == state


def test_unknown_page_is_not_a_successful_empty_search():
    scraper = rs.RealEstateScraper("Austin, TX", 500000, 3)
    assert (
        scraper.scrape(
            page_loader=lambda url: rs.HTMLFixture("<h1>Access denied</h1>"), throttle=False
        )
        == []
    )
    assert "No recognized" in scraper.last_error


def test_explicit_empty_page_has_no_failure():
    scraper = rs.RealEstateScraper("Austin, TX", 500000, 3)
    assert (
        scraper.scrape(
            page_loader=lambda url: rs.HTMLFixture('<p data-testid="no-results">None</p>'),
            throttle=False,
        )
        == []
    )
    assert scraper.last_error is None


def test_repeated_pagination_stops():
    scraper, pages = rs.demo_pages("Austin, TX", 500000, 3)
    first = scraper._build_url()
    html = pages[first].replace(scraper._build_url(2), first)
    loader = Mock(return_value=rs.HTMLFixture(html, first))
    records = scraper.scrape(page_loader=loader, throttle=False)
    assert len(records) == 9
    assert loader.call_count == 1
    assert "loop" in scraper.last_error


def test_off_site_next_link_is_not_fetched():
    scraper, pages = rs.demo_pages("Austin, TX", 500000, 3)
    first = scraper._build_url()
    html = pages[first].replace(scraper._build_url(2), "https://other.example/next")
    loader = Mock(return_value=rs.HTMLFixture(html, first))
    assert len(scraper.scrape(page_loader=loader, throttle=False)) == 9
    assert loader.call_count == 1
    assert "off-site" in scraper.last_error


def test_later_failure_retains_partial_data():
    scraper, pages = rs.demo_pages("Austin, TX", 500000, 3)
    loader = Mock(
        side_effect=[rs.HTMLFixture(pages[scraper._build_url()]), rs.PlaywrightTimeout("Timed out")]
    )
    assert len(scraper.scrape(page_loader=loader, throttle=False)) == 9
    assert "Page 2 failed" in scraper.last_error


def test_local_filters_exclude_unknown_and_out_of_range_properties():
    scraper, pages = rs.demo_pages("Austin, TX", 500000, 3, count=3)
    html = pages[scraper._build_url()]
    original = rs.generate_demo_listings("Austin, TX", 500000, 3, count=3)
    html = html.replace(original[0].price, "$900,000").replace(original[1].price, "Contact agent")
    records = scraper.scrape(page_loader=lambda url: rs.HTMLFixture(html), throttle=False)
    assert len(records) == 1
    assert scraper.stats["excluded"] == 2


def test_page_cap_does_not_load_an_extra_page():
    scraper, pages = rs.demo_pages("Austin, TX", 500000, 3)
    loader = Mock(side_effect=lambda url: rs.HTMLFixture(pages[url], url))
    assert len(scraper.scrape(1, page_loader=loader, throttle=False)) == 9
    assert loader.call_count == 1


def test_browser_closes_on_navigation_error(monkeypatch):
    browser = Mock()
    page = browser.new_context.return_value.new_page.return_value
    page.goto.side_effect = rs.PlaywrightError("Navigation failed")
    context = Mock()
    context.__enter__ = Mock(
        return_value=SimpleNamespace(chromium=SimpleNamespace(launch=lambda **kw: browser))
    )
    context.__exit__ = Mock(return_value=False)
    monkeypatch.setattr(rs, "sync_playwright", lambda: context)
    scraper = rs.RealEstateScraper("Austin, TX", 500000, 3)
    assert scraper.scrape() == []
    assert scraper.last_error
    browser.close.assert_called_once()


def test_http_error_does_not_parse_interstitial(monkeypatch):
    browser = Mock()
    page = browser.new_context.return_value.new_page.return_value
    page.goto.return_value = SimpleNamespace(status=403)
    context = Mock()
    context.__enter__ = Mock(
        return_value=SimpleNamespace(chromium=SimpleNamespace(launch=lambda **kw: browser))
    )
    context.__exit__ = Mock(return_value=False)
    monkeypatch.setattr(rs, "sync_playwright", lambda: context)
    scraper = rs.RealEstateScraper("Austin, TX", 500000, 3)
    assert scraper.scrape() == []
    assert "HTTP 403" in scraper.last_error
    page.wait_for_selector.assert_not_called()
    browser.close.assert_called_once()


@pytest.mark.parametrize(
    "args",
    [
        ["--max-pages", "0"],
        ["--max-price", "0"],
        ["--min-beds", "-1"],
        ["--city", " "],
        ["--output", "file.csv"],
    ],
)
def test_invalid_cli_arguments_fail_before_launch(args, monkeypatch):
    launch = Mock(side_effect=AssertionError("Browser must not start"))
    monkeypatch.setattr(rs, "sync_playwright", launch)
    with pytest.raises(SystemExit) as exc:
        rs.main(args)
    assert exc.value.code == 2
    launch.assert_not_called()


def test_cli_empty_result_has_failure_exit(monkeypatch):
    monkeypatch.setattr(rs.RealEstateScraper, "scrape", lambda *a, **k: [])
    assert rs.main([]) == 1


def test_cli_partial_export_is_marked(tmp_path, monkeypatch, capsys):
    def partial(self, *args):
        self.last_error = "Page 2 unavailable"
        return [rs.Listing(address="1 Oak", price="$500,000", beds="3")]

    monkeypatch.setattr(rs.RealEstateScraper, "scrape", partial)
    out = tmp_path / "partial.xlsx"
    assert rs.main(["--output", str(out)]) == 2
    assert "[PARTIAL]" in capsys.readouterr().out
    assert "PARTIAL" in openpyxl.load_workbook(out)["Export notes"]["B2"].value


def test_atomic_export_keeps_previous_file_on_failure(tmp_path, monkeypatch):
    out = tmp_path / "old.xlsx"
    out.write_bytes(b"original content")
    monkeypatch.setattr(rs.os, "replace", Mock(side_effect=PermissionError("File open")))
    with pytest.raises(PermissionError):
        rs.export_to_excel([rs.Listing(address="1 Oak")], out)
    assert out.read_bytes() == b"original content"
    assert list(tmp_path.iterdir()) == [out]


def test_excel_features_and_unknown_values():
    buffer = io.BytesIO()
    rs.export_to_excel(
        [rs.Listing(price="Contact agent", beds="Studio", baths="", sqft="1,400")], buffer
    )
    ws = openpyxl.load_workbook(buffer)["Listings"]
    assert ws["B2"].value == "Contact agent"
    assert ws["C2"].value == 0
    assert ws["D2"].value is None
    assert ws["G2"].is_date
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref == "A1:G2"


@pytest.fixture
def server():
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    yield httpd.server_port
    httpd.shutdown()
    httpd.server_close()
    thread.join()


def get(port, path, headers=None):
    conn = HTTPConnection("127.0.0.1", port, timeout=5)
    try:
        conn.request("GET", path, headers=headers or {})
        response = conn.getresponse()
        return response.status, response.read(), dict(response.getheaders())
    finally:
        conn.close()


def test_http_export_matches_filters_and_sort(server):
    status, body, _ = get(server, "/api/demo")
    assert status == 200
    assert len(json.loads(body)["listings"]) == 25
    status, body, headers = get(server, "/api/export?budget=600000&beds=3&sort=price-desc")
    assert status == 200
    assert "attachment" in headers["Content-Disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(body))["Listings"]
    records, _ = rs.run_demo(max_price=800000, min_beds=2)
    expected = filter_records(
        records, {"budget": ["600000"], "beds": ["3"], "sort": ["price-desc"]}
    )
    assert ws.max_row == len(expected) + 1
    assert [ws.cell(i + 2, 1).value for i in range(len(expected))] == [r.address for r in expected]
    assert all(ws.cell(i, 2).value <= 600000 for i in range(2, ws.max_row + 1))


@pytest.mark.parametrize(
    "path",
    [
        "/api/demo?city=Bad",
        "/api/demo?count=100000",
        "/api/export?budget=nan",
        "/api/export?beds=-2",
        "/api/export?search=zzzzmissing",
        "/api/export?sort=bad",
    ],
)
def test_invalid_api_requests_are_rejected(server, path):
    assert get(server, path)[0] == 400


def test_static_routes_and_host_restriction(server):
    for path in ["/", "/app.css", "/app.js"]:
        status, body, headers = get(server, path)
        assert status == 200 and body
        assert "Content-Security-Policy" in headers
    assert get(server, "/../realestate_scraper.py")[0] == 404
    assert get(server, "/", {"Host": "foreign.example"})[0] == 403
